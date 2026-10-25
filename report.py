"""Standardised test-report generation: PDF (ReportLab) and editable DOCX (python-docx).

Integrity model
  * `report_payload()` builds a canonical dict (instrument, conditions, observations, computed results, video
    evidence reading, attachments hashes).
  * SHA-256 of its canonical JSON is the *report hash* printed in the report and shown in the UI. The hash is
    recomputed from the database by /verify/<test-id>, so any later edit of the record is detected.
  * Optional digital signature = HMAC-SHA256(server secret, hash | signatory | time). Demo-grade; a production
    deployment would use a PKI / DSC (see README).
"""
from __future__ import annotations

import hashlib
import hmac
import io
import json
from pathlib import Path

from PIL import Image
from reportlab.graphics.barcode import qrencoder
from reportlab.graphics.shapes import Drawing, Rect
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import (Image as RLImage, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

BASE = Path(__file__).parent
IMG = BASE / "static" / "img"
NAVY, ORANGE, GREEN, RED = colors.HexColor("#0b2d5b"), colors.HexColor("#f59e0b"), colors.HexColor("#15803d"), colors.HexColor("#b91c1c")
GREY, LIGHT = colors.HexColor("#475569"), colors.HexColor("#f1f5f9")

FD = "/usr/share/fonts/truetype/dejavu/"
for name, f in (("DV", "DejaVuSans.ttf"), ("DV-B", "DejaVuSans-Bold.ttf"), ("DVS", "DejaVuSerif.ttf"),
                ("DVS-B", "DejaVuSerif-Bold.ttf"), ("DVS-I", "DejaVuSerif-Italic.ttf"), ("DVM", "DejaVuSansMono.ttf")):
    try:
        pdfmetrics.registerFont(TTFont(name, FD + f))
    except Exception:                                   # pragma: no cover - falls back to Helvetica
        pass
FONT = "DV" if "DV" in pdfmetrics.getRegisteredFontNames() else "Helvetica"
FONT_B = "DV-B" if FONT == "DV" else "Helvetica-Bold"
SERIF_B = "DVS-B" if FONT == "DV" else "Times-Bold"
SERIF_I = "DVS-I" if FONT == "DV" else "Times-Italic"


# ----------------------------------------------------------------------------- payload / hash / signature
def fmt(v, nd=3, unit=""):
    if v is None or v == "":
        return "–"
    try:
        return f"{float(v):.{nd}f}{(' ' + unit) if unit else ''}"
    except (TypeError, ValueError):
        return str(v)


def report_payload(t: dict, lab: dict, attachments: list[dict]) -> dict:
    return {
        "test_id": t["id"],
        "instrument": {k: t.get(k) for k in ("instrument_id", "instrument_type", "manufacturer", "model", "serial_number",
                                              "capacity", "division", "verification_interval", "unit", "accuracy_class",
                                              "location")},
        "standard": t.get("test_standard"), "test_date": t.get("test_date"), "uploaded_by": t.get("created_by"),
        "laboratory": lab, "conditions": t.get("env", {}), "observations": t.get("obs", {}), "results": t.get("result", {}),
        "video_evidence": {"file": t.get("video_name"), "sha256": t.get("video_sha256"),
                           "confirmed_reading_kg": t.get("reading_confirmed")},
        "attachments": [{"name": a["filename"], "sha256": a["sha256"]} for a in attachments],
    }


def payload_hash(payload: dict) -> str:
    blob = json.dumps(payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False, default=str).encode()
    return hashlib.sha256(blob).hexdigest()


def sign(secret: str, rhash: str, signer: str, ts: str) -> str:
    return hmac.new(secret.encode(), f"{rhash}|{signer}|{ts}".encode(), hashlib.sha256).hexdigest()


# ----------------------------------------------------------------------------- QR helpers (no external dependency)
def qr_matrix(data: str) -> list[list[bool]]:
    for ver in range(3, 15):
        try:
            q = qrencoder.QRCode(ver, qrencoder.QRErrorCorrectLevel.M)
            q.addData(data)
            q.make()
            n = q.getModuleCount()
            return [[bool(q.isDark(r, c)) for c in range(n)] for r in range(n)]
        except Exception:
            continue
    raise ValueError("QR payload too long")


def qr_svg(data: str, px: int = 120) -> str:
    m = qr_matrix(data)
    n = len(m)
    rects = "".join(f'<rect x="{c + 2}" y="{r + 2}" width="1" height="1"/>' for r in range(n) for c in range(n) if m[r][c])
    return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {n + 4} {n + 4}" width="{px}" height="{px}" '
            f'shape-rendering="crispEdges"><rect width="100%" height="100%" fill="#fff"/><g fill="#0b1f3a">{rects}</g></svg>')


def qr_png(data: str, scale: int = 8) -> bytes:
    m = qr_matrix(data)
    n = len(m)
    im = Image.new("RGB", ((n + 4) * scale,) * 2, "white")
    px = im.load()
    for r in range(n):
        for c in range(n):
            if m[r][c]:
                for dy in range(scale):
                    for dx in range(scale):
                        px[(c + 2) * scale + dx, (r + 2) * scale + dy] = (11, 31, 58)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def _qr_drawing(data: str, size: float) -> Drawing:
    m = qr_matrix(data)
    n = len(m)
    cell = size / (n + 2)
    d = Drawing(size, size)
    for r in range(n):
        for c in range(n):
            if m[r][c]:
                d.add(Rect((c + 1) * cell, size - (r + 2) * cell, cell, cell, fillColor=NAVY, strokeColor=None, strokeWidth=0))
    return d


# ----------------------------------------------------------------------------- PDF
class _Numbered(rl_canvas.Canvas):
    def __init__(self, *a, **kw):
        super().__init__(*a, **kw)
        self._saved = []

    def showPage(self):
        self._saved.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._saved)
        for st in self._saved:
            self.__dict__.update(st)
            self.setFont(FONT, 7.5)
            self.setFillColor(GREY)
            self.drawRightString(A4[0] - 15 * mm, 9 * mm, f"Page {self._pageNumber} of {total}")
            super().showPage()
        super().save()


def _styles():
    s = {}
    s["title"] = ParagraphStyle("title", fontName=SERIF_B, fontSize=17, alignment=TA_CENTER, textColor=NAVY, leading=21)
    s["sub"] = ParagraphStyle("sub", fontName=FONT, fontSize=9, alignment=TA_CENTER, textColor=GREY, leading=12)
    s["h"] = ParagraphStyle("h", fontName=FONT_B, fontSize=10, textColor=NAVY, spaceBefore=9, spaceAfter=3)
    s["b"] = ParagraphStyle("b", fontName=FONT, fontSize=8.5, leading=11)
    s["small"] = ParagraphStyle("small", fontName=FONT, fontSize=7.2, leading=9, textColor=GREY)
    s["mono"] = ParagraphStyle("mono", fontName="DVM" if FONT == "DV" else "Courier", fontSize=7, leading=9)
    return s


def _tbl(rows, widths, header=True, align_right_from=1, zebra=True):
    t = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
    st = [("FONT", (0, 0), (-1, -1), FONT, 8), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
          ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]
    if header:
        st += [("BACKGROUND", (0, 0), (-1, 0), NAVY), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONT", (0, 0), (-1, 0), FONT_B, 8)]
    if align_right_from is not None:
        st.append(("ALIGN", (align_right_from, 1 if header else 0), (-1, -1), "RIGHT"))
    if zebra:
        for i in range(1 if header else 0, len(rows)):
            if i % 2 == 0:
                st.append(("BACKGROUND", (0, i), (-1, i), LIGHT))
    t.setStyle(TableStyle(st))
    return t


def _kv(rows, w=(55 * mm, 125 * mm)):
    t = Table(rows, colWidths=w)
    t.setStyle(TableStyle([("FONT", (0, 0), (0, -1), FONT, 8), ("FONT", (1, 0), (1, -1), FONT_B, 8),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                           ("BACKGROUND", (0, 0), (0, -1), LIGHT), ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]))
    return t


def _badge(ok, text=None):
    text = text or ("PASS" if ok else "FAIL")
    t = Table([[text]], colWidths=[16 * mm], rowHeights=[5.5 * mm])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), GREEN if ok else RED), ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
                           ("FONT", (0, 0), (-1, -1), FONT_B, 7.5), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                           ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    return t


def _overall_color(o):
    return {"PASS": GREEN, "FAIL": RED}.get(o, ORANGE)


def build_pdf(t: dict, lab: dict, attachments: list[dict], rhash: str, verify_url: str, sig: dict | None,
              generated_at: str, attach_dir: Path | None = None) -> bytes:
    S = _styles()
    r = t["result"]
    unit = t.get("unit") or "kg"
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=34 * mm, bottomMargin=17 * mm,
                            title=f"Test Report {t['id']}", author=lab.get("lab_name", ""))

    def header(c, d):
        c.saveState()
        c.drawImage(str(IMG / "emblem.png"), 15 * mm, A4[1] - 29 * mm, width=22 * mm, height=17.6 * mm, mask="auto")
        c.drawImage(str(IMG / "logo.png"), A4[0] - 15 * mm - 62 * mm, A4[1] - 27 * mm, width=62 * mm, height=13.6 * mm, mask="auto")
        c.setFont(FONT, 7)
        c.setFillColor(GREY)
        c.drawString(15 * mm, A4[1] - 31 * mm, "Government of India")
        c.setStrokeColor(colors.HexColor("#cbd5e1"))
        c.line(15 * mm, A4[1] - 33 * mm, A4[0] - 15 * mm, A4[1] - 33 * mm)
        c.setFont(FONT, 7)
        c.drawString(15 * mm, 9 * mm, f"{t['id']}  |  {t['instrument_id']}  |  Report hash {rhash[:16]}…")
        c.restoreState()

    el = [Paragraph("TEST REPORT", S["title"]),
          Paragraph("Non-Automatic Weighing Instrument (NAWI)<br/>As per OIML R-76 &nbsp;·&nbsp; " + r.get("edition", ""), S["sub"]),
          Spacer(1, 5)]

    # 1 Instrument
    el.append(Paragraph("1. Instrument Details", S["h"]))
    el.append(_kv([["Instrument ID", t["instrument_id"]], ["Instrument Type", t["instrument_type"]],
                   ["Make / Manufacturer", t["manufacturer"]], ["Model", t["model"]],
                   ["Capacity (Max)", f"{t['capacity']:g} {unit}"], ["Division (d)", f"{t['division']:g} {unit}"],
                   ["Verification interval (e)", f"{t['verification_interval']:g} {unit}"],
                   ["Accuracy class", t["accuracy_class"]], ["Number of intervals n = Max/e", f"{r.get('n') or '–'}"],
                   ["Serial Number", t["serial_number"]], ["Location / Laboratory", t["location"]]]))
    # 2 Summary
    el.append(Paragraph("2. Test Summary", S["h"]))
    ov = r.get("overall", "–")
    badge = Table([[ov]], colWidths=[22 * mm], rowHeights=[5.8 * mm])
    badge.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), _overall_color(ov)), ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
                               ("FONT", (0, 0), (-1, -1), FONT_B, 8), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    el.append(_kv([["Test ID", t["id"]], ["Date of Test", t["test_date"]], ["Test Standard", f"{t['test_standard']} ({r.get('edition','')})"],
                   ["Overall Result", badge], ["Remarks", Paragraph(r.get("summary", ""), S["b"])]]))
    per = [["Test", "Clause", "Result"]]
    for k, tt in r.get("tests", {}).items():
        res = "PASS" if tt["pass"] else ("FAIL" if tt["evaluated"] else "NOT TESTED")
        per.append([tt["title"], tt["clause"], res])
    if len(per) > 1:
        el += [Spacer(1, 4), _tbl(per, [85 * mm, 55 * mm, 40 * mm], align_right_from=None)]

    # 3 Conditions
    env = t.get("env", {})
    el.append(Paragraph("3. Laboratory &amp; Environmental Conditions", S["h"]))
    el.append(_kv([["Laboratory", f"{lab.get('lab_name','')} – {lab.get('lab_section','')}"],
                   ["Temperature", fmt(env.get("temperature_c"), 1, "°C")], ["Relative humidity", fmt(env.get("humidity_pct"), 0, "%")],
                   ["Atmospheric pressure", fmt(env.get("pressure_hpa"), 0, "hPa")],
                   ["Standard weights used", env.get("weights_used") or "–"], ["Tested by", env.get("tested_by") or "–"],
                   ["Remarks", Paragraph(env.get("remarks") or "–", S["b"])]]))
    for i in r.get("environment_issues", []):
        el.append(Paragraph("⚠ " + i["message"], S["small"]))

    tests = r.get("tests", {})
    # 4 Weighing
    w = tests.get("weighing")
    if w and w["rows"]:
        el.append(Paragraph("4. Weighing Test (Accuracy) – " + w["clause"], S["h"]))
        rows = [[f"Load ({unit})", "MPE", "Ind. ↑", "dL ↑", "Error ↑", "OK", "Ind. ↓", "dL ↓", "Error ↓", "OK"]]
        for row in w["rows"]:
            line = [fmt(row["load"], 3), "±" + fmt(row["mpe"], 4)]
            for k in ("up", "down"):
                c = row[k]
                line += [fmt(c["indication"], 3), fmt(c["delta_l"], 4), f"{c['error']:+.4f}", "✓" if c["pass"] else "✗"] if c else ["–"] * 4
            rows.append(line)
        el.append(_tbl(rows, [21 * mm] + [17 * mm] + [17 * mm, 15 * mm, 18 * mm, 8 * mm] * 2))
        el.append(Paragraph("Error E = I + ½e − ΔL − L (rounding-corrected). Pass if |E| ≤ MPE. ↑ increasing load, ↓ decreasing load.", S["small"]))
    # 5 Ecc
    e = tests.get("eccentricity")
    if e and e["rows"]:
        el.append(Paragraph(f"5. Eccentricity Test – {e['clause']} (load {e['load']:g} {unit})", S["h"]))
        rows = [["Position", "Indication", "ΔL", "Error", "MPE", "Result"]]
        for row in e["rows"]:
            rows.append([row["position"], fmt(row["indication"], 3), fmt(row["delta_l"], 4), f"{row['error']:+.4f}", "±" + fmt(row["mpe"], 4),
                         "PASS" if row["pass"] else "FAIL"])
        el.append(_tbl(rows, [40 * mm, 30 * mm, 25 * mm, 30 * mm, 30 * mm, 25 * mm]))
    # 6 Repeat
    rp = tests.get("repeatability")
    if rp and rp["series"]:
        el.append(Paragraph(f"6. Repeatability Test – {rp['clause']}", S["h"]))
        rows = [["Load", "Errors of the weighings", "Spread", "MPE", "Result"]]
        for row in rp["series"]:
            rows.append([fmt(row["load"], 3), ",  ".join(f"{x:+.4f}" for x in row["errors"]), fmt(row["spread"], 4), "±" + fmt(row["mpe"], 4),
                         "PASS" if row["pass"] else "FAIL"])
        el.append(_tbl(rows, [25 * mm, 80 * mm, 25 * mm, 25 * mm, 25 * mm], align_right_from=2))
    # 7 Disc
    ds = tests.get("discrimination")
    if ds and ds["rows"]:
        el.append(Paragraph(f"7. Discrimination Test – {ds['clause']}", S["h"]))
        rows = [["Load", "Indication before", "Added load", "Indication after", "Change", "Required ≥ d", "Result"]]
        for row in ds["rows"]:
            rows.append([fmt(row["load"], 3), fmt(row["ind_before"], 3), fmt(row["added_load"], 4), fmt(row["ind_after"], 3),
                         fmt(row["change"], 4), fmt(row["required"], 4), "PASS" if row["pass"] else "FAIL"])
        el.append(_tbl(rows, [22 * mm, 28 * mm, 24 * mm, 28 * mm, 24 * mm, 28 * mm, 26 * mm]))

    # 8 Video evidence
    el.append(Paragraph("8. Display Reading Evidence (video)", S["h"]))
    if t.get("video_name"):
        poster = IMG / "poster.jpg"
        ev = [["Video file", t["video_name"]], ["Uploaded", t.get("video_uploaded_at") or "–"],
              ["Video SHA-256", Paragraph(t.get("video_sha256") or "(bundled sample)", S["mono"])],
              ["Reading extracted (OCR)", fmt(t.get("reading_raw"), 6, "kg")],
              ["Reading confirmed by operator", fmt(t.get("reading_confirmed"), 6, "kg")]]
        left = _kv(ev, (45 * mm, 80 * mm))
        row = [[left, RLImage(str(poster), width=44 * mm, height=31 * mm) if poster.exists() else ""]]
        el.append(Table(row, colWidths=[130 * mm, 50 * mm], style=[("VALIGN", (0, 0), (-1, -1), "TOP")]))
    else:
        el.append(Paragraph("No video evidence was attached to this test.", S["b"]))

    # 9 Attachments
    el.append(Paragraph("9. Photographs &amp; Supporting Documents", S["h"]))
    if attachments:
        el.append(_tbl([["File", "Type", "SHA-256 (first 16)"]] + [[a["filename"], a["kind"], (a["sha256"] or "")[:16]] for a in attachments],
                       [90 * mm, 30 * mm, 60 * mm], align_right_from=None))
        imgs = []
        for a in attachments:
            p = (attach_dir or Path(".")) / a["stored"]
            if a["kind"] == "photo" and p.exists():
                try:
                    with Image.open(p) as im:
                        w_, h_ = im.size
                    sc = min(55 * mm / w_, 40 * mm / h_)
                    imgs.append(RLImage(str(p), width=w_ * sc, height=h_ * sc))
                except Exception:
                    pass
        if imgs:
            rows = [imgs[i:i + 3] for i in range(0, len(imgs), 3)]
            el += [Spacer(1, 4), Table(rows, colWidths=[60 * mm] * 3, style=[("ALIGN", (0, 0), (-1, -1), "LEFT")])]
    else:
        el.append(Paragraph("None attached.", S["b"]))

    # Compliance + signature + QR
    ok_color = _overall_color(ov)
    comp = Table([[Paragraph(f"<b>COMPLIANCE RESULT: <font color='{ok_color.hexval().replace('0x', '#')}'>{ov}</font></b><br/>{r.get('summary','')}", S["b"])]],
                 colWidths=[180 * mm])
    comp.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 1, ok_color), ("BACKGROUND", (0, 0), (-1, -1), LIGHT),
                              ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    qr = _qr_drawing(verify_url, 30 * mm)
    ver = Paragraph(f"<b>Scan to verify authenticity</b><br/><font size=7>SHA-256 report hash:</font><br/><font name='{S['mono'].fontName}' size=6.6>{rhash[:32]}<br/>{rhash[32:]}</font><br/><font size=6.6 color='#475569'>{verify_url}</font>", S["b"])
    if sig:
        sig_cell = [Paragraph(f"<font name='{SERIF_I}' size=15 color='#0b2d5b'>{sig['name']}</font>", S["b"]),
                    Paragraph(f"<b>Authorized Signatory</b><br/>{sig['name']}<br/>({sig['designation']})<br/>"
                              f"<font size=6.5 color='#475569'>Digitally signed {sig['at']}<br/>HMAC {sig['mac'][:20]}…</font>", S["small"])]
    else:
        sig_cell = [Paragraph("<font color='#94a3b8'>(not digitally signed)</font>", S["b"]),
                    Paragraph(f"<b>Test Officer</b><br/>{lab.get('signatory_name','')}", S["small"])]
    foot = Table([[qr, ver, sig_cell]], colWidths=[34 * mm, 90 * mm, 56 * mm])
    foot.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEABOVE", (0, 0), (-1, 0), 0.6, colors.HexColor("#cbd5e1"))]))
    el += [Spacer(1, 8), KeepTogether([comp, Spacer(1, 6), foot,
                                       Paragraph(f"Report generated {generated_at}. This report is secured with SHA-256 hashing for authenticity and integrity.", S["small"])])]
    doc.build(el, onFirstPage=header, onLaterPages=header, canvasmaker=_Numbered)
    return buf.getvalue()


# ----------------------------------------------------------------------------- DOCX (editable)
def build_docx(t: dict, lab: dict, attachments: list[dict], rhash: str, verify_url: str, sig: dict | None,
               generated_at: str, attach_dir: Path | None = None) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    r, unit = t["result"], t.get("unit") or "kg"
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Mm(15)
        s.top_margin, s.bottom_margin = Mm(15), Mm(15)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)

    def shade(cell, hexcol):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), hexcol)
        tcPr.append(sh)

    def heading(text):
        p = d.add_paragraph()
        run = p.add_run(text)
        run.bold, run.font.size, run.font.color.rgb = True, Pt(12), RGBColor(0x0B, 0x2D, 0x5B)
        p.paragraph_format.space_before = Pt(10)

    def table(rows, header=True, widths=None):
        tb = d.add_table(rows=len(rows), cols=len(rows[0]))
        tb.style = "Table Grid"
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = tb.cell(i, j)
                c.text = ""
                run = c.paragraphs[0].add_run(str(val))
                run.font.size = Pt(9)
                if header and i == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    shade(c, "0B2D5B")
                elif not header and j == 0:
                    shade(c, "F1F5F9")
                if widths:
                    c.width = Mm(widths[j])
        return tb

    hdr = d.add_table(rows=1, cols=2)
    hdr.cell(0, 0).paragraphs[0].add_run().add_picture(str(IMG / "emblem.png"), width=Mm(22))
    hdr.cell(0, 1).paragraphs[0].add_run().add_picture(str(IMG / "logo.png"), width=Mm(62))
    hdr.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("TEST REPORT")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(20), RGBColor(0x0B, 0x2D, 0x5B)
    p = d.add_paragraph(f"Non-Automatic Weighing Instrument (NAWI)\nAs per OIML R-76 · {r.get('edition','')}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    heading("1. Instrument Details")
    table([["Instrument ID", t["instrument_id"]], ["Instrument Type", t["instrument_type"]], ["Make / Manufacturer", t["manufacturer"]],
           ["Model", t["model"]], ["Capacity (Max)", f"{t['capacity']:g} {unit}"], ["Division (d)", f"{t['division']:g} {unit}"],
           ["Verification interval (e)", f"{t['verification_interval']:g} {unit}"], ["Accuracy class", t["accuracy_class"]],
           ["n = Max/e", r.get("n") or "–"], ["Serial Number", t["serial_number"]], ["Location / Laboratory", t["location"]]],
          header=False, widths=[60, 120])
    heading("2. Test Summary")
    table([["Test ID", t["id"]], ["Date of Test", t["test_date"]], ["Test Standard", f"{t['test_standard']} ({r.get('edition','')})"],
           ["Overall Result", r.get("overall", "–")], ["Remarks", r.get("summary", "")]], header=False, widths=[60, 120])
    per = [["Test", "Clause", "Result"]] + [[x["title"], x["clause"], "PASS" if x["pass"] else ("FAIL" if x["evaluated"] else "NOT TESTED")]
                                           for x in r.get("tests", {}).values()]
    d.add_paragraph()
    table(per)

    env = t.get("env", {})
    heading("3. Laboratory & Environmental Conditions")
    table([["Laboratory", f"{lab.get('lab_name','')} – {lab.get('lab_section','')}"], ["Temperature", fmt(env.get("temperature_c"), 1, "°C")],
           ["Relative humidity", fmt(env.get("humidity_pct"), 0, "%")], ["Atmospheric pressure", fmt(env.get("pressure_hpa"), 0, "hPa")],
           ["Standard weights used", env.get("weights_used") or "–"], ["Tested by", env.get("tested_by") or "–"],
           ["Remarks", env.get("remarks") or "–"]], header=False, widths=[60, 120])

    tests = r.get("tests", {})
    w = tests.get("weighing")
    if w and w["rows"]:
        heading("4. Weighing Test (Accuracy) – " + w["clause"])
        rows = [[f"Load ({unit})", "MPE", "Ind ↑", "dL ↑", "Err ↑", "OK", "Ind ↓", "dL ↓", "Err ↓", "OK"]]
        for row in w["rows"]:
            line = [fmt(row["load"], 3), "±" + fmt(row["mpe"], 4)]
            for k in ("up", "down"):
                c = row[k]
                line += [fmt(c["indication"], 3), fmt(c["delta_l"], 4), f"{c['error']:+.4f}", "✓" if c["pass"] else "✗"] if c else ["–"] * 4
            rows.append(line)
        table(rows)
    e = tests.get("eccentricity")
    if e and e["rows"]:
        heading(f"5. Eccentricity Test – {e['clause']} (load {e['load']:g} {unit})")
        table([["Position", "Indication", "ΔL", "Error", "MPE", "Result"]] + [[x["position"], fmt(x["indication"], 3), fmt(x["delta_l"], 4),
              f"{x['error']:+.4f}", "±" + fmt(x["mpe"], 4), "PASS" if x["pass"] else "FAIL"] for x in e["rows"]])
    rp = tests.get("repeatability")
    if rp and rp["series"]:
        heading(f"6. Repeatability Test – {rp['clause']}")
        table([["Load", "Errors", "Spread", "MPE", "Result"]] + [[fmt(x["load"], 3), ", ".join(f"{v:+.4f}" for v in x["errors"]), fmt(x["spread"], 4),
              "±" + fmt(x["mpe"], 4), "PASS" if x["pass"] else "FAIL"] for x in rp["series"]])
    ds = tests.get("discrimination")
    if ds and ds["rows"]:
        heading(f"7. Discrimination Test – {ds['clause']}")
        table([["Load", "Before", "Added", "After", "Change", "Required ≥ d", "Result"]] + [[fmt(x["load"], 3), fmt(x["ind_before"], 3),
              fmt(x["added_load"], 4), fmt(x["ind_after"], 3), fmt(x["change"], 4), fmt(x["required"], 4), "PASS" if x["pass"] else "FAIL"]
              for x in ds["rows"]])

    heading("8. Display Reading Evidence (video)")
    if t.get("video_name"):
        table([["Video file", t["video_name"]], ["Uploaded", t.get("video_uploaded_at") or "–"], ["Video SHA-256", t.get("video_sha256") or "(bundled sample)"],
               ["Reading extracted (OCR)", fmt(t.get("reading_raw"), 6, "kg")], ["Reading confirmed by operator", fmt(t.get("reading_confirmed"), 6, "kg")]],
              header=False, widths=[60, 120])
        if (IMG / "poster.jpg").exists():
            d.add_picture(str(IMG / "poster.jpg"), width=Mm(55))
    else:
        d.add_paragraph("No video evidence attached.")

    heading("9. Photographs & Supporting Documents")
    if attachments:
        table([["File", "Type", "SHA-256 (first 16)"]] + [[a["filename"], a["kind"], (a["sha256"] or "")[:16]] for a in attachments])
        for a in attachments:
            pth = (attach_dir or Path(".")) / a["stored"]
            if a["kind"] == "photo" and pth.exists():
                try:
                    d.add_picture(str(pth), width=Mm(60))
                except Exception:
                    pass
    else:
        d.add_paragraph("None attached.")

    heading("Compliance Result & Verification")
    p = d.add_paragraph()
    run = p.add_run(f"{r.get('overall','–')} – ")
    run.bold = True
    p.add_run(r.get("summary", ""))
    d.add_picture(io.BytesIO(qr_png(verify_url)), width=Mm(32))
    d.add_paragraph(f"Scan to verify: {verify_url}")
    d.add_paragraph(f"SHA-256 report hash: {rhash}")
    if sig:
        d.add_paragraph(f"Digitally signed by {sig['name']} ({sig['designation']}) on {sig['at']}  ·  HMAC {sig['mac'][:20]}…")
    else:
        d.add_paragraph(f"Test Officer: {lab.get('signatory_name','')}   (not digitally signed)")
    d.add_paragraph(f"Report generated {generated_at}. This editable copy is for drafting; the PDF carries the secured original.")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()
