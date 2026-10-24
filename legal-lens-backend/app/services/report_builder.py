"""Renders the standardized NAWI test report: PDF (reportlab) and editable DOCX (python-docx).

Both renderers consume the same plain-dict context built by services/reporting.py:
  report_no, generated_at, lab_name, test_id, test_date, standard, overall_result, prepared_by, remarks,
  instrument [(label, value)], environment [(label, value)],
  observations [{label, methodology, passed, columns [[key, label]], rows [dict], warnings [str]}],
  compliance [{label, required, performed, passed}],
  attachments [{filename, kind, caption, path, sha256}], videos [{filename, sha256, uploaded_at, ...}],
  content_hash, verify_url, signature None | {name, designation, signed_at, signature}

Note: reportlab's built-in fonts have no <=, minus or Greek glyphs, so report text sticks to ASCII/Latin-1.
Drop emblem.png / logo.png into /assets to brand the header (optional).
"""
import io
from pathlib import Path
from xml.sax.saxutils import escape

from reportlab.graphics.barcode.qr import QrCodeWidget
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from reportlab.platypus import (BaseDocTemplate, Frame, Image, KeepTogether, PageTemplate,
                                Paragraph, Spacer, Table, TableStyle)

ASSETS_DIR = Path(__file__).resolve().parents[2] / "assets"
NAVY = colors.HexColor("#0B2E6B")
ORANGE = colors.HexColor("#F28C1B")
GREEN = colors.HexColor("#1B7F3B")
RED = colors.HexColor("#C62828")
GRID = colors.HexColor("#C9D1DE")
LIGHT = colors.HexColor("#F3F6FB")
IMAGE_EXT = (".png", ".jpg", ".jpeg")


def fmt(v):
    if v is None:
        return "-"
    if isinstance(v, bool):
        return "PASS" if v else "FAIL"
    if isinstance(v, float):
        if abs(v) < 1e-12:
            return "0"
        return f"{v:.6f}".rstrip("0").rstrip(".")
    return str(v)


def _res_markup(passed):
    if passed is None:
        return "-"
    color, word = ("#1B7F3B", "PASS") if passed in (True, "PASS") else ("#C62828", "FAIL" if passed in (False, "FAIL") else str(passed))
    return f'<font color="{color}"><b>{word}</b></font>'


def _overall_markup(result):
    color = {"PASS": "#1B7F3B", "FAIL": "#C62828"}.get(result, "#B26A00")
    return f'<font color="{color}"><b>{escape(str(result))}</b></font>'


class NumberedCanvas(canvas.Canvas):
    """Canvas that can print 'Page x of y' in the footer."""

    def __init__(self, *args, footer_left="", **kwargs):
        super().__init__(*args, **kwargs)
        self._footer_left = footer_left
        self._saved = []

    def showPage(self):
        self._saved.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._saved)
        for state in self._saved:
            self.__dict__.update(state)
            self._draw_footer(total)
            super().showPage()
        super().save()

    def _draw_footer(self, total):
        w, _ = A4
        self.setStrokeColor(GRID)
        self.line(15 * mm, 14 * mm, w - 15 * mm, 14 * mm)
        self.setFont("Helvetica", 7.5)
        self.setFillColor(colors.HexColor("#555555"))
        self.drawString(15 * mm, 10 * mm, self._footer_left)
        self.drawRightString(w - 15 * mm, 10 * mm, f"Page {self._pageNumber} of {total}")


def _draw_header(c, doc, ctx):
    w, h = A4
    c.saveState()
    x = 15 * mm
    emblem, logo = ASSETS_DIR / "emblem.png", ASSETS_DIR / "logo.png"
    if emblem.exists():
        c.drawImage(str(emblem), x, h - 26 * mm, width=14 * mm, height=16 * mm, preserveAspectRatio=True, mask="auto")
        x += 17 * mm
    c.setFillColor(NAVY)
    c.setFont("Helvetica-Bold", 8.5)
    c.drawString(x, h - 16 * mm, "GOVERNMENT OF INDIA")
    c.setFont("Helvetica", 7.5)
    c.drawString(x, h - 20 * mm, "Ministry of Consumer Affairs, Food & Public Distribution")
    c.drawString(x, h - 23.5 * mm, "Department of Consumer Affairs - Legal Metrology")
    if logo.exists():
        c.drawImage(str(logo), w - 15 * mm - 55 * mm, h - 26 * mm, width=55 * mm, height=16 * mm,
                    preserveAspectRatio=True, anchor="e", mask="auto")
    else:
        f, size = "Helvetica-Bold", 20
        wl, wn = c.stringWidth("LEGAL ", f, size), c.stringWidth("LENS", f, size)
        xr = w - 15 * mm - (wl + wn)
        c.setFont(f, size)
        c.setFillColor(NAVY)
        c.drawString(xr, h - 18 * mm, "LEGAL ")
        c.setFillColor(ORANGE)
        c.drawString(xr + wl, h - 18 * mm, "LENS")
        c.setFont("Helvetica", 7)
        c.setFillColor(NAVY)
        c.drawRightString(w - 15 * mm, h - 23 * mm, "TEST  |  VERIFY  |  ENSURE")
    y, seg = h - 28.5 * mm, (w - 30 * mm) / 3
    for i, col in enumerate(("#F28C1B", "#D9DEE8", "#138808")):
        c.setFillColor(colors.HexColor(col))
        c.rect(15 * mm + i * seg, y, seg, 1.2 * mm, stroke=0, fill=1)
    c.restoreState()


def _qr_drawing(value, size):
    widget = QrCodeWidget(value)
    x0, y0, x1, y1 = widget.getBounds()
    d = Drawing(size, size, transform=[size / (x1 - x0), 0, 0, size / (y1 - y0), 0, 0])
    d.add(widget)
    return d


def build_pdf(ctx: dict, path) -> None:
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=8.8, leading=11.5)
    small = ParagraphStyle("small", parent=body, fontSize=7.8, leading=10)
    cell = ParagraphStyle("cell", parent=body, fontSize=8, leading=10)
    cellh = ParagraphStyle("cellh", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)
    cellb = ParagraphStyle("cellb", parent=cell, fontName="Helvetica-Bold")
    h1 = ParagraphStyle("h1", fontName="Times-Bold", fontSize=20, leading=24, alignment=TA_CENTER, textColor=NAVY)
    sub = ParagraphStyle("sub", fontName="Times-Roman", fontSize=11, leading=14, alignment=TA_CENTER)
    subi = ParagraphStyle("subi", parent=sub, fontName="Times-Italic", fontSize=10)
    h2 = ParagraphStyle("h2", fontName="Times-Bold", fontSize=12, leading=14, textColor=NAVY, spaceBefore=10, spaceAfter=4)
    h3 = ParagraphStyle("h3", parent=h2, fontSize=10, leading=12, spaceBefore=6, spaceAfter=2)
    mono = ParagraphStyle("mono", fontName="Courier", fontSize=7.8, leading=10)

    width = A4[0] - 30 * mm
    base_style = [("GRID", (0, 0), (-1, -1), 0.4, GRID), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                  ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]

    def kv(rows, markup_keys=()):
        data = [[Paragraph(escape(str(k)), cellb), Paragraph(v if k in markup_keys else escape(fmt(v) if v != "" else "-"), cell)]
                for k, v in rows]
        t = Table(data, colWidths=[width * 0.32, width * 0.68])
        t.setStyle(TableStyle(base_style + [("BACKGROUND", (0, 0), (0, -1), LIGHT)]))
        return t

    def grid(columns, rows, widths=None):
        n = len(columns)
        data = [[Paragraph(escape(lbl), cellh) for _, lbl in columns]]
        for r in rows:
            data.append([Paragraph(_res_markup(r.get(k)) if k == "passed" else escape(fmt(r.get(k))), cell) for k, _ in columns])
        t = Table(data, colWidths=widths or [width / n] * n, repeatRows=1)
        t.setStyle(TableStyle(base_style + [("BACKGROUND", (0, 0), (-1, 0), NAVY)]))
        return t

    story = [Spacer(1, 2 * mm), Paragraph("TEST REPORT", h1),
             Paragraph("Non-Automatic Weighing Instrument (NAWI)", sub),
             Paragraph(f"As per {escape(ctx['standard'])}", subi), Spacer(1, 3 * mm)]

    story += [Paragraph("1. Instrument Details", h2), kv(ctx["instrument"])]
    story += [Paragraph("2. Test Summary", h2),
              kv([("Report No.", ctx["report_no"]), ("Test ID", ctx["test_id"]), ("Date of Test", ctx["test_date"]),
                  ("Test Standard", ctx["standard"]), ("Laboratory", ctx["lab_name"]),
                  ("Prepared by", ctx["prepared_by"]), ("Report generated on", ctx["generated_at"]),
                  ("Overall Result", _overall_markup(ctx["overall_result"]))], markup_keys=("Overall Result",))]
    story += [Paragraph("3. Laboratory and Environmental Conditions", h2), kv(ctx["environment"])]

    story.append(Paragraph("4. Test Observations and Calculations", h2))
    for i, ob in enumerate(ctx["observations"], 1):
        block = [Paragraph(f"4.{i} {escape(ob['label'])} - {_res_markup(ob['passed'])}", h3),
                 Paragraph(f"<i>Method: {escape(ob['methodology'])}</i>", small), Spacer(1, 2), grid(ob["columns"], ob["rows"])]
        if ob.get("warnings"):
            block.append(Paragraph("Notes: " + escape("; ".join(ob["warnings"])), small))
        story.append(KeepTogether(block))
    if not ctx["observations"]:
        story.append(Paragraph("No observations recorded.", body))

    story.append(Paragraph("5. Compliance Summary", h2))
    comp_rows = [{"label": c["label"], "required": "Yes" if c["required"] else "Optional",
                  "performed": "Yes" if c["performed"] else "No", "passed": c["passed"]} for c in ctx["compliance"]]
    story.append(grid([["label", "Test"], ["required", "Required"], ["performed", "Performed"], ["passed", "Result"]],
                      comp_rows, widths=[width * 0.5, width * 0.17, width * 0.17, width * 0.16]))
    story.append(Spacer(1, 2 * mm))
    story.append(Paragraph(f"Overall compliance with {escape(ctx['standard'])}: {_overall_markup(ctx['overall_result'])}", body))
    if ctx.get("remarks"):
        story += [Paragraph("Remarks", h3), Paragraph(escape(ctx["remarks"]), body)]

    if ctx["attachments"] or ctx["videos"]:
        story.append(Paragraph("6. Attachments and Evidence", h2))
    for a in ctx["attachments"]:
        story.append(Paragraph(f"<b>{escape(a['filename'])}</b> ({escape(a['kind'])}) {escape(a.get('caption') or '')} "
                               f"<font size=6.5>SHA-256 {escape(a['sha256'][:16])}...</font>", small))
        if a["filename"].lower().endswith(IMAGE_EXT) and Path(a["path"]).exists():
            try:
                iw, ih = ImageReader(a["path"]).getSize()
                scale = min(80 * mm / iw, 60 * mm / ih, 1.0 if iw < 80 * mm else 80 * mm / iw)
                story += [Spacer(1, 1.5 * mm), Image(a["path"], width=iw * scale, height=ih * scale, hAlign="LEFT"), Spacer(1, 2 * mm)]
            except Exception:
                pass
    if ctx["videos"]:
        vrows = [{"file": v["filename"], "reading": v.get("confirmed_reading"), "auto": v.get("extracted_reading"),
                  "method": v.get("method") or "-", "sha": v["sha256"][:16] + "..."} for v in ctx["videos"]]
        story += [Spacer(1, 2 * mm), Paragraph("Test video evidence", h3),
                  grid([["file", "Video file"], ["reading", "Confirmed reading"], ["auto", "Auto-detected"],
                        ["method", "Method"], ["sha", "SHA-256 (start)"]], vrows)]

    qr = _qr_drawing(ctx["verify_url"], 30 * mm)
    hash_lines = "<br/>".join(ctx["content_hash"][i:i + 32] for i in (0, 32))
    mid = [Paragraph("<b>Report verification</b>", body),
           Paragraph("Scan the QR code or open the link to verify the authenticity of this report.", small), Spacer(1, 2),
           Paragraph("<b>SHA-256 (report content)</b>", small), Paragraph(hash_lines, mono),
           Paragraph(escape(ctx["verify_url"]), small)]
    sig = ctx.get("signature")
    if sig:
        right = [Paragraph("<b>Digitally signed by</b>", small), Paragraph(escape(sig["name"]), body),
                 Paragraph(escape(sig["designation"]), small), Paragraph(f"Signed on: {escape(sig['signed_at'])}", small),
                 Paragraph(f"Signature ID: {escape(sig['signature'][:20])}...", small)]
    else:
        right = [Paragraph("<b>Authorised Signatory</b>", small), Spacer(1, 10 * mm), Paragraph("______________________", small),
                 Paragraph("Not yet signed (digital signature is optional)", small)]
    box = Table([[qr, mid, right]], colWidths=[34 * mm, 86 * mm, width - 120 * mm])
    box.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, NAVY), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                             ("LINEAFTER", (0, 0), (1, 0), 0.3, GRID), ("BACKGROUND", (0, 0), (-1, -1), LIGHT),
                             ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story += [Spacer(1, 6 * mm), KeepTogether([Paragraph("7. Verification and Authorisation", h2), box])]

    doc = BaseDocTemplate(str(path), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=34 * mm,
                          bottomMargin=18 * mm, title=f"Test Report {ctx['report_no']}", author="Legal Lens",
                          subject=f"NAWI test report - {ctx['test_id']}")
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main",
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc.addPageTemplates([PageTemplate(id="main", frames=[frame], onPage=lambda c, d: _draw_header(c, d, ctx))])
    footer = f"Report No: {ctx['report_no']}  |  Test ID: {ctx['test_id']}  |  SHA-256: {ctx['content_hash'][:16]}..."
    doc.build(story, canvasmaker=lambda *a, **k: NumberedCanvas(*a, footer_left=footer, **k))


# ----------------------------------------------------------------------------- DOCX
def build_docx(ctx: dict, path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(15)
    sec.top_margin, sec.bottom_margin = Mm(18), Mm(18)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(9.5)
    navy, green, red = RGBColor(0x0B, 0x2E, 0x6B), RGBColor(0x1B, 0x7F, 0x3B), RGBColor(0xC6, 0x28, 0x28)

    def shade(cell, fill):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(shd)

    def write(cell, text, bold=False, color=None, size=9):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold, run.font.size = bold, Pt(size)
        if text in ("PASS", "FAIL"):
            run.bold, run.font.color.rgb = True, (green if text == "PASS" else red)
        elif color is not None:
            run.font.color.rgb = color

    def heading(text, size=13):
        p = doc.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(3)
        r = p.add_run(text)
        r.bold, r.font.size, r.font.color.rgb, r.font.name = True, Pt(size), navy, "Times New Roman"

    def kv(rows):
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for k, v in rows:
            cells = t.add_row().cells
            write(cells[0], str(k), bold=True)
            shade(cells[0], "F3F6FB")
            write(cells[1], fmt(v) if v != "" else "-")
            cells[0].width, cells[1].width = Mm(55), Mm(125)

    def grid(columns, rows):
        t = doc.add_table(rows=1, cols=len(columns))
        t.style = "Table Grid"
        for i, (_, lbl) in enumerate(columns):
            write(t.rows[0].cells[i], lbl, bold=True, color=RGBColor(255, 255, 255), size=8.5)
            shade(t.rows[0].cells[i], "0B2E6B")
        for r in rows:
            cells = t.add_row().cells
            for i, (k, _) in enumerate(columns):
                write(cells[i], fmt(r.get(k)), size=8.5)

    # header / footer
    hp = sec.header.paragraphs[0]
    hr = hp.add_run("GOVERNMENT OF INDIA | Department of Consumer Affairs        ")
    hr.font.size, hr.font.color.rgb = Pt(8), navy
    r1 = hp.add_run("LEGAL ")
    r1.bold, r1.font.size, r1.font.color.rgb = True, Pt(16), navy
    r2 = hp.add_run("LENS")
    r2.bold, r2.font.size, r2.font.color.rgb = True, Pt(16), RGBColor(0xF2, 0x8C, 0x1B)
    fp = sec.footer.paragraphs[0]
    fr = fp.add_run(f"Report No: {ctx['report_no']} | Test ID: {ctx['test_id']} | SHA-256: {ctx['content_hash'][:16]}...   Page ")
    fr.font.size = Pt(7.5)
    for instr in ("PAGE",):
        run = fp.add_run()
        for typ, text in (("begin", None), (None, instr), ("end", None)):
            if typ:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), typ)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)
        run.font.size = Pt(7.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("TEST REPORT")
    tr.bold, tr.font.size, tr.font.color.rgb, tr.font.name = True, Pt(22), navy, "Times New Roman"
    for text, italic in (("Non-Automatic Weighing Instrument (NAWI)", False), (f"As per {ctx['standard']}", True)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.italic, run.font.size = italic, Pt(11)

    heading("1. Instrument Details")
    kv(ctx["instrument"])
    heading("2. Test Summary")
    kv([("Report No.", ctx["report_no"]), ("Test ID", ctx["test_id"]), ("Date of Test", ctx["test_date"]),
        ("Test Standard", ctx["standard"]), ("Laboratory", ctx["lab_name"]), ("Prepared by", ctx["prepared_by"]),
        ("Report generated on", ctx["generated_at"]), ("Overall Result", ctx["overall_result"])])
    heading("3. Laboratory and Environmental Conditions")
    kv(ctx["environment"])
    heading("4. Test Observations and Calculations")
    for i, ob in enumerate(ctx["observations"], 1):
        heading(f"4.{i} {ob['label']} - {'PASS' if ob['passed'] else 'FAIL'}", size=10.5)
        p = doc.add_paragraph()
        run = p.add_run("Method: " + ob["methodology"])
        run.italic, run.font.size = True, Pt(8)
        grid(ob["columns"], ob["rows"])
        if ob.get("warnings"):
            p = doc.add_paragraph()
            p.add_run("Notes: " + "; ".join(ob["warnings"])).font.size = Pt(8)
    heading("5. Compliance Summary")
    grid([["label", "Test"], ["required", "Required"], ["performed", "Performed"], ["passed", "Result"]],
         [{"label": c["label"], "required": "Yes" if c["required"] else "Optional",
           "performed": "Yes" if c["performed"] else "No", "passed": c["passed"]} for c in ctx["compliance"]])
    p = doc.add_paragraph()
    p.add_run(f"\nOverall compliance with {ctx['standard']}: ").bold = True
    run = p.add_run(ctx["overall_result"])
    run.bold, run.font.color.rgb = True, (green if ctx["overall_result"] == "PASS" else red)
    if ctx.get("remarks"):
        heading("Remarks", size=10.5)
        doc.add_paragraph(ctx["remarks"])

    if ctx["attachments"] or ctx["videos"]:
        heading("6. Attachments and Evidence")
    for a in ctx["attachments"]:
        doc.add_paragraph(f"{a['filename']} ({a['kind']}) {a.get('caption') or ''} - SHA-256 {a['sha256'][:16]}...")
        if a["filename"].lower().endswith(IMAGE_EXT) and Path(a["path"]).exists():
            try:
                doc.add_picture(a["path"], width=Mm(80))
            except Exception:
                pass
    if ctx["videos"]:
        heading("Test video evidence", size=10.5)
        grid([["file", "Video file"], ["reading", "Confirmed reading"], ["auto", "Auto-detected"], ["method", "Method"],
              ["sha", "SHA-256 (start)"]],
             [{"file": v["filename"], "reading": v.get("confirmed_reading"), "auto": v.get("extracted_reading"),
               "method": v.get("method") or "-", "sha": v["sha256"][:16] + "..."} for v in ctx["videos"]])

    heading("7. Verification and Authorisation")
    try:
        import qrcode  # optional dependency (in requirements.txt)
        buf = io.BytesIO()
        qrcode.make(ctx["verify_url"]).save(buf, format="PNG")
        buf.seek(0)
        doc.add_picture(buf, width=Mm(28))
    except Exception:
        pass
    doc.add_paragraph(f"Verification URL: {ctx['verify_url']}")
    doc.add_paragraph(f"SHA-256 (report content): {ctx['content_hash']}")
    sig = ctx.get("signature")
    if sig:
        doc.add_paragraph(f"Digitally signed by {sig['name']} ({sig['designation']}) on {sig['signed_at']}. "
                          f"Signature ID: {sig['signature'][:20]}...")
    else:
        doc.add_paragraph("Authorised Signatory: ______________________   (not yet signed; digital signature is optional)")
    doc.save(str(path))
